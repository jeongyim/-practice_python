#5_1barChart.py
from openpyxl import load_workbook
wb=load_workbook("sample.xlsx")
ws=wb.active
# Reference:차트를 생성할 범위를 지정해주기 위해서 임포트
from openpyxl.chart import BarChart, Reference
#B2:D5까지의 데이터를 차트로 생성
bar_value=Reference(ws, min_row=2,max_row=5,min_col=2,max_col=4)
bar_chart=BarChart()#차트종류 설정(Bar, Line,Pie..)
bar_chart.add_data(bar_value)#차트에 데이터 추가
ws.add_chart(bar_chart, "E1")#차트를 “E1”위치에 넣음
wb.save("sample_BarChart.xlsx")
wb.close()