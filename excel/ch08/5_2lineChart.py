
from openpyxl import load_workbook
wb=load_workbook("sample.xlsx")
ws=wb.active
from openpyxl.chart import LineChart, Reference
#계열1..대신 과목명을 넣기 위해 1라인도 범위에 넣어 준다.
line_value=Reference(ws,min_row=1,max_row=5,min_col=2,max_col=4)
line_chart=LineChart()#사용할 차트종류
line_chart.add_data(line_value,titles_from_data=True)#과목명 사용할 것임
line_chart.title="성적표"
line_chart.style=10 #미리 정의된 스타일을 적용
line_chart.y_axis.title = "점수" #y축의 제목
line_chart.x_axis.title = "번호" #x축의 제목
ws.add_chart(line_chart, "E1")

wb.save("sample_lineChart.xlsx")
wb.close()
